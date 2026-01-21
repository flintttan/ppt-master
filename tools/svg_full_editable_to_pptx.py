#!/usr/bin/env python3
"""Convert an SVG-image-based PPTX into a fully editable PPTX (shapes + text).

This tool reads embedded per-slide SVGs from an input PPTX (ppt/media/imageN.svg)
then recreates each slide using editable PowerPoint objects:
- SVG primitives (rect/line/circle/polygon/path) -> PPT shapes
- SVG <text> -> PPT textboxes (editable)
- Speaker notes are copied from the original PPTX

Design goal: maximize editability in WPS/Office. Fidelity is best-effort:
- SVG opacity is approximated by blending with white (python-pptx doesn't expose alpha)
- SVG gradients are approximated by using the first stop color
- SVG curves/arcs are flattened into polyline segments for PPT freeform shapes

Usage:
  . .venv-ppt/bin/activate
  python ppt-master/tools/svg_full_editable_to_pptx.py \
    FabriqueAIHub_答辩_ppt169_20260119/FabriqueAIHub_答辩.pptx \
    -o FabriqueAIHub_答辩_ppt169_20260119/FabriqueAIHub_答辩_全可编辑版.pptx
"""

from __future__ import annotations

import argparse
import math
import re
import subprocess
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

EMU_PER_INCH = 914400
EMU_PER_PIXEL = EMU_PER_INCH / 96.0
SVG_NS = "http://www.w3.org/2000/svg"


def _local_name(tag: str) -> str:
    return tag.split("}", 1)[1] if "}" in tag else tag


def _parse_float(value: str | None, default: float = 0.0) -> float:
    if value is None:
        return default
    try:
        return float(value)
    except Exception:
        return default


def _parse_points(points: str) -> list[tuple[float, float]]:
    pts: list[tuple[float, float]] = []
    for token in re.split(r"\s+", points.strip()):
        if not token:
            continue
        if "," not in token:
            continue
        x_s, y_s = token.split(",", 1)
        pts.append((_parse_float(x_s), _parse_float(y_s)))
    return pts


def _parse_color_hex(color: str | None, default: tuple[int, int, int] = (0, 0, 0)) -> tuple[int, int, int]:
    if not color:
        return default
    color = color.strip()
    if color.startswith("url("):
        # handled separately (gradient) by caller
        return default
    if not color.startswith("#"):
        return default
    hexv = color[1:]
    if len(hexv) == 3:
        try:
            return (int(hexv[0] * 2, 16), int(hexv[1] * 2, 16), int(hexv[2] * 2, 16))
        except Exception:
            return default
    if len(hexv) == 6:
        try:
            return (int(hexv[0:2], 16), int(hexv[2:4], 16), int(hexv[4:6], 16))
        except Exception:
            return default
    return default


def _blend_with_white(rgb: tuple[int, int, int], opacity: float) -> tuple[int, int, int]:
    # PPT shape fill/line colors don't expose alpha in python-pptx.
    o = min(1.0, max(0.0, opacity))
    r, g, b = rgb
    return (
        int(round(r * o + 255 * (1.0 - o))),
        int(round(g * o + 255 * (1.0 - o))),
        int(round(b * o + 255 * (1.0 - o))),
    )


def _rgb_from_svg_color(color: str | None, opacity: float) -> RGBColor:
    rgb = _parse_color_hex(color, default=(0, 0, 0))
    r2, g2, b2 = _blend_with_white(rgb, opacity)
    return RGBColor(r2, g2, b2)


def _first_font_name(font_family: str | None) -> str | None:
    if not font_family:
        return None
    first = font_family.split(",", 1)[0].strip().strip('"\'')
    return first or None


def _is_cjk(s: str) -> bool:
    for ch in s:
        o = ord(ch)
        if 0x4E00 <= o <= 0x9FFF:
            return True
        if 0x3400 <= o <= 0x4DBF:
            return True
        if 0xF900 <= o <= 0xFAFF:
            return True
    return False


def _estimate_text_width_px(text: str, font_px: float) -> float:
    width_em = 0.0
    for ch in text:
        o = ord(ch)
        if ch.isspace():
            width_em += 0.35
        elif 0x4E00 <= o <= 0x9FFF or 0x3400 <= o <= 0x4DBF or 0xF900 <= o <= 0xFAFF:
            width_em += 1.0
        elif ch in "•·-–—|丨/\\:：,，.。;；()（）[]【】{}“”\"'" :
            width_em += 0.35
        else:
            width_em += 0.55
    return max(4.0, width_em * font_px + 6.0)


def _px_to_emu(px: float) -> int:
    return int(round(px * EMU_PER_PIXEL))


def _px_to_pt(px: float) -> float:
    return px * 72.0 / 96.0


def _ppt_align_from_anchor(anchor: str) -> int:
    a = anchor.lower()
    if a == "middle":
        return PP_ALIGN.CENTER
    if a == "end":
        return PP_ALIGN.RIGHT
    return PP_ALIGN.LEFT


def _bold_from_weight(font_weight: str | None) -> bool:
    if not font_weight:
        return False
    w = font_weight.strip().lower()
    if w == "bold":
        return True
    m = re.match(r"^(\d+)$", w)
    if m:
        try:
            return int(m.group(1)) >= 600
        except Exception:
            return False
    return False


# 2D affine matrix a,c,e / b,d,f
Matrix = tuple[float, float, float, float, float, float]


def mat_identity() -> Matrix:
    return (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)


def mat_mul(m2: Matrix, m1: Matrix) -> Matrix:
    # Composition m = m2 ∘ m1 (apply m1 then m2)
    a2, b2, c2, d2, e2, f2 = m2
    a1, b1, c1, d1, e1, f1 = m1
    return (
        a2 * a1 + c2 * b1,
        b2 * a1 + d2 * b1,
        a2 * c1 + c2 * d1,
        b2 * c1 + d2 * d1,
        a2 * e1 + c2 * f1 + e2,
        b2 * e1 + d2 * f1 + f2,
    )


def mat_apply(m: Matrix, x: float, y: float) -> tuple[float, float]:
    a, b, c, d, e, f = m
    return (a * x + c * y + e, b * x + d * y + f)


def parse_transform(transform: str | None) -> Matrix:
    if not transform:
        return mat_identity()

    s = transform.strip()
    m = mat_identity()
    for func, args in re.findall(r"([a-zA-Z]+)\(([^\)]*)\)", s):
        parts = [p for p in re.split(r"[ ,]+", args.strip()) if p]
        vals = [_parse_float(p, 0.0) for p in parts]
        func = func.lower()
        if func == "translate":
            tx = vals[0] if len(vals) >= 1 else 0.0
            ty = vals[1] if len(vals) >= 2 else 0.0
            t = (1.0, 0.0, 0.0, 1.0, tx, ty)
            m = mat_mul(m, t)
        elif func == "scale":
            sx = vals[0] if len(vals) >= 1 else 1.0
            sy = vals[1] if len(vals) >= 2 else sx
            t = (sx, 0.0, 0.0, sy, 0.0, 0.0)
            m = mat_mul(m, t)
        else:
            continue
    return m


@dataclass
class Style:
    fill: str | None = None
    fill_opacity: float = 1.0
    stroke: str | None = None
    stroke_opacity: float = 1.0
    stroke_width: float = 0.0
    stroke_dasharray: str | None = None


def merge_style(parent: Style, el: ET.Element) -> Style:
    s = Style(
        fill=parent.fill,
        fill_opacity=parent.fill_opacity,
        stroke=parent.stroke,
        stroke_opacity=parent.stroke_opacity,
        stroke_width=parent.stroke_width,
        stroke_dasharray=parent.stroke_dasharray,
    )

    if "fill" in el.attrib:
        s.fill = el.attrib.get("fill")
    if "fill-opacity" in el.attrib:
        s.fill_opacity = _parse_float(el.attrib.get("fill-opacity"), s.fill_opacity)
    if "stroke" in el.attrib:
        s.stroke = el.attrib.get("stroke")
    if "stroke-opacity" in el.attrib:
        s.stroke_opacity = _parse_float(el.attrib.get("stroke-opacity"), s.stroke_opacity)
    if "stroke-width" in el.attrib:
        s.stroke_width = _parse_float(el.attrib.get("stroke-width"), s.stroke_width)
    if "stroke-dasharray" in el.attrib:
        s.stroke_dasharray = el.attrib.get("stroke-dasharray")

    # Some elements use generic opacity.
    if "opacity" in el.attrib and "fill-opacity" not in el.attrib:
        s.fill_opacity = _parse_float(el.attrib.get("opacity"), s.fill_opacity)

    return s


def _extract_gradient_stop_color(svg_root: ET.Element, url_fill: str) -> tuple[str | None, float]:
    # url_fill format: url(#id)
    m = re.match(r"url\(#([^\)]+)\)", url_fill.strip())
    if not m:
        return (None, 1.0)
    grad_id = m.group(1)

    # Find first stop.
    for el in svg_root.iter():
        if el.attrib.get("id") != grad_id:
            continue
        for stop in el.findall(f".//{{{SVG_NS}}}stop"):
            col = stop.attrib.get("stop-color")
            op = _parse_float(stop.attrib.get("stop-opacity"), 1.0)
            if col:
                return (col, op)
    return (None, 1.0)


# --- SVG path parsing & flattening ---


def _tokenize_path(d: str) -> list[str]:
    # Commands or numbers (including scientific notation)
    return re.findall(r"[AaCcHhLlMmVvZz]|[-+]?(?:\d*\.\d+|\d+)(?:[eE][-+]?\d+)?", d)


def _cubic_bezier(p0, p1, p2, p3, t: float) -> tuple[float, float]:
    x0, y0 = p0
    x1, y1 = p1
    x2, y2 = p2
    x3, y3 = p3
    u = 1.0 - t
    x = (
        u * u * u * x0
        + 3 * u * u * t * x1
        + 3 * u * t * t * x2
        + t * t * t * x3
    )
    y = (
        u * u * u * y0
        + 3 * u * u * t * y1
        + 3 * u * t * t * y2
        + t * t * t * y3
    )
    return (x, y)


def _angle(u: tuple[float, float], v: tuple[float, float]) -> float:
    ux, uy = u
    vx, vy = v
    dot = ux * vx + uy * vy
    det = ux * vy - uy * vx
    return math.atan2(det, dot)


def _arc_to_points(
    p0: tuple[float, float],
    rx: float,
    ry: float,
    phi_deg: float,
    large_arc: int,
    sweep: int,
    p1: tuple[float, float],
    max_seg_len: float = 6.0,
) -> list[tuple[float, float]]:
    # Implementation based on SVG arc implementation notes.
    x1, y1 = p0
    x2, y2 = p1
    if rx == 0.0 or ry == 0.0:
        return [p1]

    rx = abs(rx)
    ry = abs(ry)

    phi = math.radians(phi_deg % 360.0)
    cos_phi = math.cos(phi)
    sin_phi = math.sin(phi)

    dx = (x1 - x2) / 2.0
    dy = (y1 - y2) / 2.0

    x1p = cos_phi * dx + sin_phi * dy
    y1p = -sin_phi * dx + cos_phi * dy

    # Correct radii
    lam = (x1p * x1p) / (rx * rx) + (y1p * y1p) / (ry * ry)
    if lam > 1.0:
        s = math.sqrt(lam)
        rx *= s
        ry *= s

    # Center
    num = rx * rx * ry * ry - rx * rx * y1p * y1p - ry * ry * x1p * x1p
    den = rx * rx * y1p * y1p + ry * ry * x1p * x1p
    if den == 0.0:
        coef = 0.0
    else:
        coef = math.sqrt(max(0.0, num / den))
        if large_arc == sweep:
            coef = -coef

    cxp = coef * (rx * y1p / ry)
    cyp = coef * (-ry * x1p / rx)

    cx = cos_phi * cxp - sin_phi * cyp + (x1 + x2) / 2.0
    cy = sin_phi * cxp + cos_phi * cyp + (y1 + y2) / 2.0

    # Angles
    ux = (x1p - cxp) / rx
    uy = (y1p - cyp) / ry
    vx = (-x1p - cxp) / rx
    vy = (-y1p - cyp) / ry

    theta1 = math.atan2(uy, ux)
    dtheta = _angle((ux, uy), (vx, vy))

    if sweep == 0 and dtheta > 0:
        dtheta -= 2.0 * math.pi
    elif sweep == 1 and dtheta < 0:
        dtheta += 2.0 * math.pi

    # Determine segment count based on arc length.
    # Approx arc length ~ avg radius * abs(dtheta)
    avg_r = (rx + ry) / 2.0
    arc_len = avg_r * abs(dtheta)
    segs = max(4, int(math.ceil(arc_len / max_seg_len)))
    segs = min(segs, 96)

    pts: list[tuple[float, float]] = []
    for i in range(1, segs + 1):
        t = i / segs
        theta = theta1 + dtheta * t
        xep = rx * math.cos(theta)
        yep = ry * math.sin(theta)
        x = cos_phi * xep - sin_phi * yep + cx
        y = sin_phi * xep + cos_phi * yep + cy
        pts.append((x, y))
    return pts


@dataclass
class Subpath:
    points: list[tuple[float, float]]
    closed: bool


def try_parse_rounded_rect_path(d: str) -> tuple[float, float, float, float, float] | None:
    toks = _tokenize_path(d)
    if not toks:
        return None
    if any(t.isalpha() and t.islower() for t in toks if re.match(r"^[A-Za-z]$", t)):
        return None

    i = 0

    def expect_cmd(c: str) -> bool:
        nonlocal i
        if i >= len(toks) or toks[i] != c:
            return False
        i += 1
        return True

    def next_num() -> float:
        nonlocal i
        v = float(toks[i])
        i += 1
        return v

    if not expect_cmd("M"):
        return None
    if i + 1 >= len(toks):
        return None
    x0 = next_num(); y0 = next_num()

    if not expect_cmd("H"):
        return None
    x1 = next_num()

    if not expect_cmd("A"):
        return None
    rx1 = next_num(); ry1 = next_num()
    phi1 = next_num(); laf1 = int(next_num()); sf1 = int(next_num())
    x2 = next_num(); y2 = next_num()

    if not expect_cmd("V"):
        return None
    y3 = next_num()

    if not expect_cmd("A"):
        return None
    rx2 = next_num(); ry2 = next_num()
    phi2 = next_num(); laf2 = int(next_num()); sf2 = int(next_num())
    x4 = next_num(); y4 = next_num()

    if not expect_cmd("H"):
        return None
    x5 = next_num()

    if not expect_cmd("A"):
        return None
    rx3 = next_num(); ry3 = next_num()
    phi3 = next_num(); laf3 = int(next_num()); sf3 = int(next_num())
    x6 = next_num(); y6 = next_num()

    if not expect_cmd("V"):
        return None
    y7 = next_num()

    if not expect_cmd("A"):
        return None
    rx4 = next_num(); ry4 = next_num()
    phi4 = next_num(); laf4 = int(next_num()); sf4 = int(next_num())
    x8 = next_num(); y8 = next_num()

    if not expect_cmd("Z"):
        return None

    if i != len(toks):
        return None

    rxs = [rx1, rx2, rx3, rx4]
    rys = [ry1, ry2, ry3, ry4]
    if not all(abs(rx - rxs[0]) < 1e-6 for rx in rxs) or not all(abs(ry - rys[0]) < 1e-6 for ry in rys):
        return None
    if rxs[0] <= 0.0 or rys[0] <= 0.0:
        return None

    flags = [(laf1, sf1), (laf2, sf2), (laf3, sf3), (laf4, sf4)]
    if not all(laf == 0 and sf == 1 for (laf, sf) in flags):
        return None
    if not all(abs(phi) < 1e-6 for phi in (phi1, phi2, phi3, phi4)):
        return None

    left = x6
    top = y0
    right = x2
    bottom = y4
    if right <= left or bottom <= top:
        return None

    if abs(x8 - x0) > 1e-6 or abs(y8 - y0) > 1e-6:
        return None

    return (left, top, right - left, bottom - top, rxs[0])


def flatten_path(d: str) -> list[Subpath]:
    toks = _tokenize_path(d)
    i = 0
    cmd = None
    cur = (0.0, 0.0)
    start = (0.0, 0.0)
    subpaths: list[Subpath] = []
    pts: list[tuple[float, float]] = []
    closed = False

    def flush():
        nonlocal pts, closed
        if pts:
            subpaths.append(Subpath(points=pts, closed=closed))
        pts = []
        closed = False

    while i < len(toks):
        t = toks[i]
        if re.match(r"^[A-Za-z]$", t):
            cmd = t
            i += 1
            if cmd.upper() == "Z":
                closed = True
                cur = start
                flush()
                cmd = None
            continue
        if cmd is None:
            # Invalid path
            break

        is_rel = cmd.islower()
        c = cmd.upper()

        def next_num() -> float:
            nonlocal i
            v = float(toks[i])
            i += 1
            return v

        if c == "M":
            x = next_num()
            y = next_num()
            if is_rel:
                x += cur[0]
                y += cur[1]
            flush()
            cur = (x, y)
            start = cur
            pts = [cur]
            # Subsequent coordinate pairs are treated as implicit L
            cmd = "l" if is_rel else "L"

        elif c == "L":
            x = next_num()
            y = next_num()
            if is_rel:
                x += cur[0]
                y += cur[1]
            cur = (x, y)
            pts.append(cur)

        elif c == "H":
            x = next_num()
            if is_rel:
                x += cur[0]
            cur = (x, cur[1])
            pts.append(cur)

        elif c == "V":
            y = next_num()
            if is_rel:
                y += cur[1]
            cur = (cur[0], y)
            pts.append(cur)

        elif c == "C":
            x1 = next_num(); y1 = next_num()
            x2 = next_num(); y2 = next_num()
            x = next_num(); y = next_num()
            if is_rel:
                x1 += cur[0]; y1 += cur[1]
                x2 += cur[0]; y2 += cur[1]
                x += cur[0]; y += cur[1]
            p0 = cur
            p1 = (x1, y1)
            p2 = (x2, y2)
            p3 = (x, y)
            # Sample curve
            steps = 12
            for s in range(1, steps + 1):
                pts.append(_cubic_bezier(p0, p1, p2, p3, s / steps))
            cur = p3

        elif c == "A":
            rx = next_num(); ry = next_num()
            phi = next_num()
            large_arc = int(next_num())
            sweep = int(next_num())
            x = next_num(); y = next_num()
            if is_rel:
                x += cur[0]
                y += cur[1]
            p1 = (x, y)
            pts.extend(_arc_to_points(cur, rx, ry, phi, large_arc, sweep, p1))
            cur = p1

        else:
            # Unsupported command
            break

    flush()
    return subpaths


def extract_notes_text(pptx_zip: zipfile.ZipFile, slide_num: int) -> str | None:
    notes_path = f"ppt/notesSlides/notesSlide{slide_num}.xml"
    try:
        raw = pptx_zip.read(notes_path)
    except KeyError:
        return None

    try:
        root = ET.fromstring(raw)
    except Exception:
        return None

    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    paras: list[str] = []
    for p in root.findall(".//a:p", ns):
        parts = [t.text or "" for t in p.findall(".//a:t", ns)]
        s = "".join(parts).strip()
        if s:
            paras.append(s)
    return "\n".join(paras) if paras else None


def apply_shape_style(shape, style: Style, svg_root: ET.Element) -> None:
    # Fill
    fill = style.fill
    fill_op = style.fill_opacity

    if fill_op <= 0.001:
        fill = None

    if fill is None or fill == "none":
        try:
            shape.fill.background()
        except Exception:
            pass
    else:
        if fill.startswith("url("):
            col, op = _extract_gradient_stop_color(svg_root, fill)
            fill = col or "#FFFFFF"
            fill_op *= op
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb_from_svg_color(fill, fill_op)

    # Line
    stroke = style.stroke
    stroke_op = style.stroke_opacity

    if stroke_op <= 0.001:
        stroke = None

    if not stroke or stroke == "none" or style.stroke_width <= 0.0:
        try:
            shape.line.fill.background()
        except Exception:
            pass
        return

    if stroke.startswith("url("):
        col, op = _extract_gradient_stop_color(svg_root, stroke)
        stroke = col or "#000000"
        stroke_op *= op

    shape.line.color.rgb = _rgb_from_svg_color(stroke, stroke_op)
    shape.line.width = Pt(max(0.25, _px_to_pt(style.stroke_width)))

    if style.stroke_dasharray and style.stroke_dasharray.strip().lower() != "none":
        # Very small mapping: any dasharray -> dashed.
        try:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE

            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        except Exception:
            pass


def add_text(slide, svg_text: ET.Element, style: Style, transform: Matrix) -> None:
    raw = "".join(svg_text.itertext()).strip("\n")
    if raw.strip() == "":
        return

    x = _parse_float(svg_text.attrib.get("x"), 0.0)
    y = _parse_float(svg_text.attrib.get("y"), 0.0)
    anchor = (svg_text.attrib.get("text-anchor") or "start").strip()

    font_family = svg_text.attrib.get("font-family")
    font_px = _parse_float(svg_text.attrib.get("font-size"), 16.0)
    font_weight = svg_text.attrib.get("font-weight")
    fill = svg_text.attrib.get("fill") or style.fill or "#000000"
    fill_op = _parse_float(svg_text.attrib.get("fill-opacity"), style.fill_opacity)

    # Apply transform
    x, y = mat_apply(transform, x, y)

    font_pt = max(6.0, _px_to_pt(font_px))
    top_px = y - 0.8 * font_px

    est_w_px = _estimate_text_width_px(raw, font_px)
    if anchor == "middle":
        left_px = x - est_w_px / 2.0
    elif anchor == "end":
        left_px = x - est_w_px
    else:
        left_px = x

    left_px = max(0.0, min(1280.0 - 2.0, left_px))
    top_px = max(0.0, min(720.0 - 2.0, top_px))
    w_px = min(1280.0 - left_px, max(est_w_px, 20.0))
    h_px = max(font_px * 1.5, 10.0)

    shape = slide.shapes.add_textbox(
        Emu(_px_to_emu(left_px)),
        Emu(_px_to_emu(top_px)),
        Emu(_px_to_emu(w_px)),
        Emu(_px_to_emu(h_px)),
    )

    try:
        shape.fill.background()
        shape.line.fill.background()
    except Exception:
        pass

    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    try:
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
    except Exception:
        pass

    p = tf.paragraphs[0]
    p.alignment = _ppt_align_from_anchor(anchor)
    run = p.add_run()
    run.text = raw

    font = run.font
    font.size = Pt(font_pt)

    name = _first_font_name(font_family)
    if not name:
        name = "Source Han Sans SC" if _is_cjk(raw) else "Montserrat"
    font.name = name
    font.color.rgb = _rgb_from_svg_color(fill, fill_op)
    font.bold = _bold_from_weight(font_weight)


def _set_rounded_rect_adjustment(shape, rx_px: float, width_px: float, height_px: float) -> None:
    try:
        adjs = shape.adjustments
    except Exception:
        return

    denom = min(width_px, height_px)
    if rx_px <= 0.0 or denom <= 0.0 or len(adjs) < 1:
        return

    adj = (2.0 * rx_px) / denom
    if adj < 0.0:
        adj = 0.0
    if adj > 0.5:
        adj = 0.5

    try:
        adjs[0] = adj
    except Exception:
        pass


def add_rect(slide, el: ET.Element, style: Style, transform: Matrix, svg_root: ET.Element) -> None:
    x = _parse_float(el.attrib.get("x"), 0.0)
    y = _parse_float(el.attrib.get("y"), 0.0)
    w = _parse_float(el.attrib.get("width"), 0.0)
    h = _parse_float(el.attrib.get("height"), 0.0)
    rx = _parse_float(el.attrib.get("rx"), 0.0)

    (x1, y1) = mat_apply(transform, x, y)
    (x2, y2) = mat_apply(transform, x + w, y + h)
    left = min(x1, x2)
    top = min(y1, y2)
    width = abs(x2 - x1)
    height = abs(y2 - y1)

    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rx > 0.0 else MSO_SHAPE.RECTANGLE

    shp = slide.shapes.add_shape(
        shape_type,
        Emu(_px_to_emu(left)),
        Emu(_px_to_emu(top)),
        Emu(_px_to_emu(width)),
        Emu(_px_to_emu(height)),
    )
    if rx > 0.0:
        _set_rounded_rect_adjustment(shp, rx, width, height)
    apply_shape_style(shp, style, svg_root)


def add_circle(slide, el: ET.Element, style: Style, transform: Matrix, svg_root: ET.Element) -> None:
    cx = _parse_float(el.attrib.get("cx"), 0.0)
    cy = _parse_float(el.attrib.get("cy"), 0.0)
    r = _parse_float(el.attrib.get("r"), 0.0)

    x = cx - r
    y = cy - r
    w = 2 * r
    h = 2 * r

    (x1, y1) = mat_apply(transform, x, y)
    (x2, y2) = mat_apply(transform, x + w, y + h)
    left = min(x1, x2)
    top = min(y1, y2)
    width = abs(x2 - x1)
    height = abs(y2 - y1)

    shp = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Emu(_px_to_emu(left)),
        Emu(_px_to_emu(top)),
        Emu(_px_to_emu(width)),
        Emu(_px_to_emu(height)),
    )
    apply_shape_style(shp, style, svg_root)


def add_line(slide, el: ET.Element, style: Style, transform: Matrix, svg_root: ET.Element) -> None:
    x1 = _parse_float(el.attrib.get("x1"), 0.0)
    y1 = _parse_float(el.attrib.get("y1"), 0.0)
    x2 = _parse_float(el.attrib.get("x2"), 0.0)
    y2 = _parse_float(el.attrib.get("y2"), 0.0)

    (sx, sy) = mat_apply(transform, x1, y1)
    (ex, ey) = mat_apply(transform, x2, y2)

    shp = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Emu(_px_to_emu(sx)),
        Emu(_px_to_emu(sy)),
        Emu(_px_to_emu(ex)),
        Emu(_px_to_emu(ey)),
    )
    # Lines use stroke attributes.
    apply_shape_style(shp, style, svg_root)


def add_polygon(slide, el: ET.Element, style: Style, transform: Matrix, svg_root: ET.Element) -> None:
    pts = _parse_points(el.attrib.get("points", ""))
    if len(pts) < 2:
        return
    pts = [mat_apply(transform, x, y) for x, y in pts]
    pts_emu = [(_px_to_emu(x), _px_to_emu(y)) for (x, y) in pts]

    (sx, sy) = pts_emu[0]
    fb = slide.shapes.build_freeform(start_x=sx, start_y=sy)
    fb.add_line_segments(pts_emu[1:], close=True)
    shp = fb.convert_to_shape()
    apply_shape_style(shp, style, svg_root)


def add_path(slide, el: ET.Element, style: Style, transform: Matrix, svg_root: ET.Element) -> None:
    d = el.attrib.get("d")
    if not d:
        return

    rr = try_parse_rounded_rect_path(d)
    if rr is not None:
        left, top, width, height, rx = rr
        (x1, y1) = mat_apply(transform, left, top)
        (x2, y2) = mat_apply(transform, left + width, top + height)
        s_left = min(x1, x2)
        s_top = min(y1, y2)
        s_w = abs(x2 - x1)
        s_h = abs(y2 - y1)
        shp = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(_px_to_emu(s_left)),
            Emu(_px_to_emu(s_top)),
            Emu(_px_to_emu(s_w)),
            Emu(_px_to_emu(s_h)),
        )
        rx2, _ = mat_apply(transform, left + rx, top)
        rx_px = abs(rx2 - x1)
        _set_rounded_rect_adjustment(shp, rx_px, s_w, s_h)
        apply_shape_style(shp, style, svg_root)
        return

    subpaths = flatten_path(d)
    if not subpaths:
        return

    for sp in subpaths:
        pts = [mat_apply(transform, x, y) for (x, y) in sp.points]
        if len(pts) < 2:
            continue
        pts_emu = [(_px_to_emu(x), _px_to_emu(y)) for (x, y) in pts]
        (sx, sy) = pts_emu[0]
        fb = slide.shapes.build_freeform(start_x=sx, start_y=sy)
        fb.add_line_segments(pts_emu[1:], close=sp.closed or (style.fill not in (None, "none")))
        shp = fb.convert_to_shape()
        apply_shape_style(shp, style, svg_root)


def iter_svg_children(el: ET.Element) -> Iterator[ET.Element]:
    for child in list(el):
        yield child


def process_svg(slide, svg_root: ET.Element, icons_dir: Path) -> None:
    def walk(node: ET.Element, parent_style: Style, parent_tf: Matrix):
        tag = _local_name(node.tag)
        tf = mat_mul(parse_transform(node.attrib.get("transform")), parent_tf)
        st = merge_style(parent_style, node)

        if tag in ("defs", "title", "desc"):
            return

        # Handle icon placeholders (<use data-icon=.../>)
        if tag == "use" and "data-icon" in node.attrib:
            icon_name = node.attrib.get("data-icon")
            # fallback mapping
            if icon_name == "pencil-ruler":
                icon_name = "ruler"
            x = _parse_float(node.attrib.get("x"), 0.0)
            y = _parse_float(node.attrib.get("y"), 0.0)
            w = _parse_float(node.attrib.get("width"), 16.0)
            h = _parse_float(node.attrib.get("height"), 16.0)
            fill = node.attrib.get("fill") or st.fill

            icon_path = icons_dir / f"{icon_name}.svg"
            if icon_path.exists():
                icon_svg = ET.fromstring(icon_path.read_text(encoding="utf-8"))
                # icons are designed in a 16x16 box
                scale = w / 16.0 if w else 1.0
                icon_tf = mat_mul((scale, 0.0, 0.0, scale, x, y), tf)
                icon_style = Style(fill=fill, fill_opacity=st.fill_opacity)
                for child in list(icon_svg):
                    # icons usually contain paths
                    walk(child, icon_style, icon_tf)
            return

        if tag == "g":
            for child in iter_svg_children(node):
                walk(child, st, tf)
            return

        if tag == "rect":
            add_rect(slide, node, st, tf, svg_root)
        elif tag == "circle":
            add_circle(slide, node, st, tf, svg_root)
        elif tag == "line":
            add_line(slide, node, st, tf, svg_root)
        elif tag == "polygon":
            add_polygon(slide, node, st, tf, svg_root)
        elif tag == "path":
            add_path(slide, node, st, tf, svg_root)
        elif tag == "text":
            add_text(slide, node, st, tf)
        else:
            # Unknown element: recurse if it has children.
            for child in iter_svg_children(node):
                walk(child, st, tf)

    walk(svg_root, Style(fill=None), mat_identity())


def build_pptx(input_pptx: Path, output_pptx: Path, max_slides: int | None) -> None:
    icons_dir = Path(__file__).parent.parent / "templates" / "icons"

    with zipfile.ZipFile(input_pptx) as z:
        # Determine slide count from embedded SVGs.
        svg_members: list[tuple[int, str]] = []
        for i in range(1, 500):
            name = f"ppt/media/image{i}.svg"
            try:
                z.getinfo(name)
            except KeyError:
                break
            svg_members.append((i, name))

        if max_slides is not None:
            svg_members = svg_members[:max_slides]

        prs = Presentation()
        prs.slide_width = Emu(_px_to_emu(1280))
        prs.slide_height = Emu(_px_to_emu(720))
        blank = prs.slide_layouts[6]

        for slide_num, member in svg_members:
            svg_bytes = z.read(member)
            svg_root = ET.fromstring(svg_bytes)

            slide = prs.slides.add_slide(blank)

            notes = extract_notes_text(z, slide_num)
            if notes:
                try:
                    slide.notes_slide.notes_text_frame.text = notes
                except Exception:
                    pass

            process_svg(slide, svg_root, icons_dir)

    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_pptx))


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert SVG media slides into fully editable PPTX")
    parser.add_argument("input_pptx", type=Path)
    parser.add_argument("-o", "--output", type=Path, required=True)
    parser.add_argument("--max-slides", type=int, default=None, help="limit for quick test")
    args = parser.parse_args()

    if not args.input_pptx.exists():
        raise SystemExit(f"Input not found: {args.input_pptx}")

    build_pptx(args.input_pptx, args.output, args.max_slides)


if __name__ == "__main__":
    main()
